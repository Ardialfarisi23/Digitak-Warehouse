import openpyxl

wb = openpyxl.load_workbook(r'D:\Project\digitak-warehouse\data import.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)
ws = wb['Rekap BoQ Periode 2026']
print('Max row:', ws.max_row, 'Max col:', ws.max_column)
print('Row 2 (Cluster IDs):', [ws.cell(row=2, column=c).value for c in range(1, 70)])
print('Row 3 (Project Names):', [ws.cell(row=3, column=c).value for c in range(1, 70)])
print('Row 4 (Homepass Plan):', [ws.cell(row=4, column=c).value for c in range(1, 70)])
print('Row 5 (Tiket Release):', [ws.cell(row=5, column=c).value for c in range(1, 70)])
print('Row 8 (Headers B-D):', [ws.cell(row=8, column=c).value for c in range(2, 5)])
print('Row 9 (Material B-D):', [ws.cell(row=9, column=c).value for c in range(2, 5)])
print('Row 64 (Material B-D):', [ws.cell(row=64, column=c).value for c in range(2, 5)])

# Print column headers for clusters
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=8, column=c).value
    if val:
        print(f'  Col {c} (row 8): {val}')

# Print first few cluster blocks (row 2)
print('\nCluster blocks (row 2):')
for c in range(10, 60):
    val = ws.cell(row=2, column=c).value
    if val:
        print(f'  Col {c}: {val}')